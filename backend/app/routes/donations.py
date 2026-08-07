# backend/app/routes/donations.py

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from sqlalchemy import or_
from sqlalchemy.orm import Session, joinedload
from datetime import datetime
from typing import Optional

from app.database import SessionLocal
from app import models, schemas
from app.security import require_admin, require_collector, require_superadmin
from app.websocket_manager import manager
from app.routes.finance import write_audit, write_ledger, invalidate_dashboard_cache
from app.utils.payment_ledger import generate_receipt_id
from app.utils.payment_notify import notify_donation_payment
from app.rate_limit import rate_limit

router = APIRouter(prefix="/donations", tags=["Donations"])

VALID_DONOR_TYPES = {"app_user", "member", "walk_in", "anonymous"}


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def _build_donation_out(d: models.Donation) -> schemas.DonationOut:
    fund_name = d.fund_rel.name if d.fund_rel else None
    return schemas.DonationOut(
        id=d.id,
        donor_name=d.donor_name,
        amount=d.amount,
        method=d.method,
        note=d.note,
        recorded_by=d.recorded_by,
        created_at=d.created_at,
        head_id=d.head_id,
        receipt_id=d.receipt_id,
        purpose_id=d.purpose_id,
        fund_id=d.fund_id,
        fund_name=fund_name,
        donor_type=d.donor_type,
        phone=d.phone,
        chanda_no=d.chanda_no,
        donation_date=d.donation_date,
        receipt_image=d.receipt_image,
    )


def _base_query(db: Session):
    return (
        db.query(models.Donation)
        .options(
            joinedload(models.Donation.fund_rel),
            joinedload(models.Donation.purpose_rel),
        )
    )


def _validate_donation_payload(db: Session, data: schemas.DonationCreate) -> None:
    if data.amount is not None and data.amount <= 0:
        raise HTTPException(400, "Amount must be greater than zero")

    if data.donor_type is not None:
        donor_type = data.donor_type.strip().lower()
        if donor_type not in VALID_DONOR_TYPES:
            raise HTTPException(400, "Invalid donor type")

    if data.member_id is not None:
        head = db.query(models.ApprovedHead).filter_by(id=data.member_id).first()
        if not head:
            raise HTTPException(404, "Member not found")
        if not head.is_active:
            raise HTTPException(400, "Member is inactive")

    if data.purpose_id is not None:
        purpose_obj = db.query(models.DonationPurpose).filter_by(
            id=data.purpose_id, is_active=True, is_archived=False
        ).first()
        if not purpose_obj:
            raise HTTPException(404, "Donation purpose not found or inactive")

    if data.fund_id is not None:
        fund_obj = db.query(models.Fund).filter_by(id=data.fund_id, is_archived=False).first()
        if not fund_obj:
            raise HTTPException(404, "Fund not found or archived")
        if fund_obj.status == "archived":
            raise HTTPException(400, "Cannot use archived or completed funds for new donations")


# ─────────────────────────────────────────────
# GET ALL DONATIONS
# ─────────────────────────────────────────────
@router.get("/", response_model=list[schemas.DonationOut])
def get_donations(
    search: Optional[str] = None,
    method: Optional[str] = None,
    fund_id: Optional[int] = None,
    donor_type: Optional[str] = None,
    purpose_id: Optional[int] = None,
    from_date: Optional[str] = Query(None),
    to_date: Optional[str] = Query(None),
    limit: int = Query(100, ge=1, le=500),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    query = _base_query(db)

    if search:
        term = f"%{search}%"
        query = query.filter(or_(
            models.Donation.donor_name.ilike(term),
            models.Donation.receipt_id.ilike(term),
            models.Donation.phone.ilike(term),
            models.Donation.chanda_no.ilike(term),
        ))
    if method:
        query = query.filter(models.Donation.method == method)
    if fund_id is not None:
        query = query.filter(models.Donation.fund_id == fund_id)
    if donor_type:
        query = query.filter(models.Donation.donor_type == donor_type)
    if purpose_id is not None:
        query = query.filter(models.Donation.purpose_id == purpose_id)
    if from_date:
        query = query.filter(models.Donation.created_at >= datetime.fromisoformat(from_date))
    if to_date:
        query = query.filter(models.Donation.created_at <= datetime.fromisoformat(to_date))

    rows = query.order_by(models.Donation.created_at.desc()).offset(offset).limit(limit).all()
    return [_build_donation_out(d) for d in rows]


# ─────────────────────────────────────────────
# SUMMARY
# ─────────────────────────────────────────────
@router.get("/summary")
def donation_summary(
    fund_id: Optional[int] = None,
    purpose_id: Optional[int] = None,
    donor_type: Optional[str] = None,
    from_date: Optional[str] = Query(None),
    to_date: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    query = db.query(models.Donation)
    if fund_id is not None:
        query = query.filter(models.Donation.fund_id == fund_id)
    if purpose_id is not None:
        query = query.filter(models.Donation.purpose_id == purpose_id)
    if donor_type:
        query = query.filter(models.Donation.donor_type == donor_type)
    if from_date:
        query = query.filter(models.Donation.created_at >= datetime.fromisoformat(from_date))
    if to_date:
        query = query.filter(models.Donation.created_at <= datetime.fromisoformat(to_date))

    donations = query.all()
    total_amount = sum(d.amount for d in donations)
    total_count = len(donations)

    by_method: dict[str, float] = {}
    for d in donations:
        m = (d.method or "cash").strip().lower()
        by_method[m] = round(by_method.get(m, 0) + d.amount, 2)

    return {
        "total_donations": total_count,
        "total_amount": round(total_amount, 2),
        "cash_total": by_method.get("cash", 0.0),
        "upi_total": by_method.get("upi", 0.0),
        "bank_total": by_method.get("bank", 0.0),
        "cheque_total": by_method.get("cheque", 0.0),
        "other_total": sum(v for k, v in by_method.items() if k not in {"cash", "upi", "bank", "cheque"}),
    }


# ─────────────────────────────────────────────
# ADD DONATION
# ─────────────────────────────────────────────
@router.post("/", response_model=schemas.DonationOut, status_code=201)
async def add_donation(
    data: schemas.DonationCreate,
    request: Request,
    db: Session = Depends(get_db),
    user: dict = Depends(require_collector),
):
    await rate_limit(request, "collector_write")
    _validate_donation_payload(db, data)
    actor_id = int(user.get("sub"))
    db_user  = db.query(models.User).filter_by(id=actor_id).first()
    recorded_by = db_user.name if db_user else "Admin"

    head = None
    if data.member_id:
        head = db.query(models.ApprovedHead).filter_by(id=data.member_id).first()

    donor_name = (data.donor_name or "").strip() or (head.name if head else "")
    if not donor_name:
        raise HTTPException(400, "Donor name is required")

    chanda_no = data.chanda_no
    if not chanda_no and head:
        chanda_no = head.chanda_no

    donation = models.Donation(
        donor_name=donor_name,
        user_id=actor_id,
        collector_id=actor_id,
        amount=data.amount,
        method=(data.method or "cash").lower(),
        note=(data.note or "").strip() or None,
        recorded_by=recorded_by,
        created_at=datetime.utcnow(),
        head_id=head.id if head else None,
        purpose_id=data.purpose_id,
        fund_id=data.fund_id,
        donor_type=(data.donor_type or "walk_in").strip().lower(),
        phone=data.phone or (head.phone if head else None),
        chanda_no=chanda_no,
        donation_date=data.donation_date,
        receipt_image=data.receipt_image,
        receipt_id=generate_receipt_id(db, "DN", datetime.utcnow()),
    )

    db.add(donation)
    db.flush()

    write_ledger(
        db, "donation", "income", data.amount,
        family_id=head.id if head else None,
        donation_id=donation.id,
        note=f"Donation by {donation.donor_name}",
        created_by_id=actor_id,
    )
    write_audit(
        db, "donations", donation.id, "create",
        new_values={
            "amount": data.amount,
            "method": data.method,
            "donor_name": donation.donor_name,
            "fund_id": data.fund_id,
            "donor_type": data.donor_type,
        },
        performed_by_id=actor_id,
    )

    db.commit()
    db.refresh(donation)

    invalidate_dashboard_cache()

    # Reload with relationships
    donation = _base_query(db).filter(models.Donation.id == donation.id).first()
    manager.publish_sync("finance", "donation_created", {
        "donation_id": donation.id,
        "head_id": donation.head_id,
        "fund_id": donation.fund_id,
    })
    manager.publish_sync("finance", "dashboard_updated", {})
    notify_donation_payment(db, donation)

    return _build_donation_out(donation)


# ── Donation Report (PDF / Excel) ─────────────────────────────────────────────
# MUST be above /{donation_id} — otherwise FastAPI matches "report" as an int
# parameter and returns 422 Unprocessable Content.

@router.get("/report")
def download_donation_report(
    format: str = "pdf",           # "pdf" | "excel"
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    fund_id: Optional[int] = None,
    method: Optional[str] = None,
    purpose_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: dict = Depends(require_admin),
):
    import io as _io

    from_dt = None
    to_dt   = None
    try:
        if from_date: from_dt = datetime.fromisoformat(from_date)
        if to_date:   to_dt   = datetime.fromisoformat(to_date)
    except ValueError:
        pass

    q = db.query(models.Donation)
    if from_dt:   q = q.filter(models.Donation.created_at >= from_dt)
    if to_dt:     q = q.filter(models.Donation.created_at <= to_dt)
    if fund_id:   q = q.filter(models.Donation.fund_id == fund_id)
    if method:    q = q.filter(models.Donation.method == method)
    if purpose_id: q = q.filter(models.Donation.purpose_id == purpose_id)
    donations = q.order_by(models.Donation.created_at.desc()).all()
    total = round(sum(float(d.amount or 0) for d in donations), 2)

    rows = [
        {
            "receipt_id":  d.receipt_id or "—",
            "donor":       d.donor_name or "Unknown",
            "purpose":     (d.purpose_rel.name if d.purpose_rel else "—"),
            "fund":        (d.fund_rel.name if d.fund_rel else "—"),
            "method":      (d.method or "—").upper(),
            "source":      d.donor_type or "walk_in",
            "date":        d.created_at.strftime("%d %b %Y") if d.created_at else "—",
            "amount":      float(d.amount or 0),
        }
        for d in donations
    ]

    if format == "excel":
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise HTTPException(500, "openpyxl not installed")
        from fastapi.responses import StreamingResponse

        wb  = openpyxl.Workbook()
        ws  = wb.active
        ws.title = "Donations"

        hdr_fill = PatternFill("solid", fgColor="0F5C4C")
        hdr_font = Font(bold=True, color="FFFFFF", size=11)

        header = ["Receipt ID", "Donor", "Purpose", "Fund", "Method", "Source", "Date", "Amount (₹)"]
        col_w  = [20, 22, 20, 18, 10, 12, 14, 14]
        for col, (h, w) in enumerate(zip(header, col_w), 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

        for i, r in enumerate(rows, 2):
            ws.append([r["receipt_id"], r["donor"], r["purpose"], r["fund"],
                       r["method"], r["source"], r["date"], r["amount"]])
            ws.cell(i, 8).number_format = '#,##0.00'

        total_row = len(rows) + 2
        ws.cell(total_row, 7, "TOTAL").font = Font(bold=True)
        ws.cell(total_row, 8, total).font   = Font(bold=True)
        ws.cell(total_row, 8).number_format = '#,##0.00'

        buf = _io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": "attachment; filename=donations_report.xlsx"})

    # ── PDF ───────────────────────────────────────────────────────────────────
    try:
        from reportlab.lib import colors as rl_colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        raise HTTPException(500, "reportlab not installed")
    from fastapi.responses import StreamingResponse

    buf = _io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5*cm, rightMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    elems  = []

    period = ""
    if from_date: period += f" from {from_date}"
    if to_date:   period += f" to {to_date}"
    elems.append(Paragraph(f"Donations Report{period}", styles["Title"]))
    elems.append(Paragraph(f"Total: ₹{total:,.2f}  |  Count: {len(rows)}", styles["Normal"]))
    elems.append(Spacer(1, 0.4*cm))

    tdata = [["Receipt ID", "Donor", "Purpose", "Fund", "Method", "Date", "Amount (₹)"]]
    for r in rows:
        tdata.append([r["receipt_id"], r["donor"], r["purpose"], r["fund"],
                      r["method"], r["date"], f"₹{r['amount']:,.0f}"])
    tdata.append(["", "", "", "", "", "TOTAL", f"₹{total:,.0f}"])

    col_ws = [3.5*cm, 4*cm, 3.5*cm, 3*cm, 2*cm, 3*cm, 2.5*cm]
    t = Table(tdata, colWidths=col_ws, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND",   (0, 0), (-1, 0), rl_colors.HexColor("#0F5C4C")),
        ("TEXTCOLOR",    (0, 0), (-1, 0), rl_colors.white),
        ("FONTNAME",     (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",     (0, 0), (-1, 0), 9),
        ("FONTSIZE",     (0, 1), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [rl_colors.white, rl_colors.HexColor("#F7F9F8")]),
        ("FONTNAME",     (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID",         (0, 0), (-1, -1), 0.25, rl_colors.HexColor("#DDDDDD")),
        ("ALIGN",        (-1, 0), (-1, -1), "RIGHT"),
        ("VALIGN",       (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING",   (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING",(0, 0), (-1, -1), 4),
    ]))
    elems.append(t)
    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": "attachment; filename=donations_report.pdf"})


# ─────────────────────────────────────────────
# GET SINGLE
# ─────────────────────────────────────────────
@router.get("/{donation_id}", response_model=schemas.DonationOut)
def get_single_donation(
    donation_id: int,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    donation = _base_query(db).filter(models.Donation.id == donation_id).first()
    if not donation:
        raise HTTPException(404, "Donation not found")
    return _build_donation_out(donation)


# ─────────────────────────────────────────────
# UPDATE
# ─────────────────────────────────────────────
@router.put("/{donation_id}", response_model=schemas.DonationOut)
def update_donation(
    donation_id: int,
    data: schemas.DonationCreate,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    donation = db.query(models.Donation).filter(models.Donation.id == donation_id).first()
    if not donation:
        raise HTTPException(404, "Donation not found")

    _validate_donation_payload(db, data)

    actor_id = int(user.get("sub"))
    old = {"amount": donation.amount, "method": donation.method, "fund_id": donation.fund_id}
    old_amount = donation.amount
    old_fund_id = donation.fund_id

    head = None
    if data.member_id is not None:
        head = db.query(models.ApprovedHead).filter_by(id=data.member_id).first()
        if not head:
            raise HTTPException(404, "Member not found")
        donation.head_id = head.id
        donation.chanda_no = data.chanda_no or head.chanda_no
        donation.phone = data.phone or head.phone
        donation.donor_name = head.name

    if data.donor_name is not None:
        donor_name = data.donor_name.strip()
        if donor_name:
            donation.donor_name = donor_name

    if data.amount is not None:
        donation.amount = data.amount
    if data.method is not None:
        donation.method = data.method.lower()
    if data.note is not None:
        donation.note = (data.note or "").strip() or None
    if data.donor_type is not None:
        donor_type = data.donor_type.strip().lower()
        if donor_type not in VALID_DONOR_TYPES:
            raise HTTPException(400, "Invalid donor type")
        donation.donor_type = donor_type
    if data.phone is not None:
        donation.phone = data.phone
    if data.chanda_no is not None:
        donation.chanda_no = data.chanda_no
    if data.donation_date is not None:
        donation.donation_date = data.donation_date
    if data.receipt_image is not None:
        donation.receipt_image = data.receipt_image
    if data.purpose_id is not None:
        donation.purpose_id = data.purpose_id
    if data.fund_id is not None and data.fund_id != donation.fund_id:
        donation.fund_id = data.fund_id

    if data.amount is not None and data.amount != old_amount:
        diff = data.amount - old_amount
        write_ledger(
            db, "adjustment", "expense" if diff < 0 else "income",
            abs(diff),
            donation_id=donation.id,
            note=f"Donation amount adjusted from {old_amount:.2f} to {data.amount:.2f}",
            created_by_id=actor_id,
        )

    if data.fund_id is not None and data.fund_id != old_fund_id:
        write_ledger(
            db, "adjustment", "expense",
            abs(donation.amount),
            donation_id=donation.id,
            note=f"Donation fund changed from {old_fund_id} to {data.fund_id}",
            created_by_id=actor_id,
        )

    write_audit(
        db, "donations", donation.id, "update",
        old_values=old,
        new_values={"amount": donation.amount, "method": donation.method, "fund_id": donation.fund_id},
        performed_by_id=actor_id,
    )

    db.commit()
    invalidate_dashboard_cache()
    donation = _base_query(db).filter(models.Donation.id == donation_id).first()
    manager.publish_sync("finance", "donation_updated", {"donation_id": donation.id})
    manager.publish_sync("finance", "dashboard_updated", {})
    return _build_donation_out(donation)


# ─────────────────────────────────────────────
# DELETE
# ─────────────────────────────────────────────
@router.delete("/{donation_id}")
def delete_donation(
    donation_id: int,
    db: Session = Depends(get_db),
    user: dict = Depends(require_superadmin),
):
    donation = db.query(models.Donation).filter(models.Donation.id == donation_id).first()
    if not donation:
        raise HTTPException(404, "Donation not found")

    actor_id = int(user.get("sub"))
    write_ledger(
        db, "adjustment", "expense",
        donation.amount,
        donation_id=donation.id,
        note=f"Donation deleted (receipt {donation.receipt_id})",
        created_by_id=actor_id,
    )
    write_audit(
        db, "donations", donation.id, "delete",
        old_values={"amount": donation.amount, "donor_name": donation.donor_name},
        performed_by_id=actor_id,
    )
    db.delete(donation)
    db.commit()
    invalidate_dashboard_cache()
    manager.publish_sync("finance", "donation_deleted", {"donation_id": donation_id})
    manager.publish_sync("finance", "dashboard_updated", {})

    return {"message": "Donation deleted"}

