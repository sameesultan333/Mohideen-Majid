# backend/app/routes/funds.py
# Finance Fund (Campaign) system

from __future__ import annotations

import io
import math
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload

from app import models, schemas
from app.database import SessionLocal
from app.routes.finance import write_audit, invalidate_dashboard_cache
from app.security import require_admin, require_superadmin, get_current_user
from app.websocket_manager import manager

router = APIRouter(prefix="/funds", tags=["Funds"])

VALID_STATUSES = {"active", "archived"}
# Legacy statuses mapped to active for backward compat
_LEGACY_STATUS_MAP = {"draft": "active", "completed": "active"}


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def _actor(user: dict, db: Session) -> tuple[int, str]:
    uid = int(user.get("sub"))
    u = db.query(models.User).filter_by(id=uid).first()
    return uid, (u.name if u else "Admin")


def _get_fund(db: Session, fund_id: int) -> models.Fund:
    fund = db.query(models.Fund).filter_by(id=fund_id).first()
    if not fund:
        raise HTTPException(404, "Fund not found")
    return fund


def _fund_stats(db: Session, fund_id: int) -> schemas.FundStats:
    donation_stats = (
        db.query(
            func.count(models.Donation.id),
            func.coalesce(func.sum(models.Donation.amount), 0.0),
            func.count(func.distinct(models.Donation.phone)),
            func.max(models.Donation.created_at),
        )
        .filter(models.Donation.fund_id == fund_id)
        .one()
    )
    expense_stats = (
        db.query(
            func.count(models.Expense.id),
            func.coalesce(func.sum(models.Expense.amount), 0.0),
        )
        .filter(
            models.Expense.fund_id == fund_id,
            models.Expense.is_deleted == False,
            models.Expense.approved_at != None,   # only approved expenses count
        )
        .one()
    )

    total_collected = round(float(donation_stats[1] or 0.0), 2)
    total_spent = round(float(expense_stats[1] or 0.0), 2)
    balance = round(total_collected - total_spent, 2)

    fund = db.query(models.Fund).filter_by(id=fund_id).first()
    goal = fund.goal_amount if fund else None
    pct = round(total_collected / goal * 100, 1) if goal and goal > 0 else None

    return schemas.FundStats(
        total_donations=int(donation_stats[0] or 0),
        total_collected=total_collected,
        total_expenses=int(expense_stats[0] or 0),
        total_spent=total_spent,
        balance=balance,
        goal_amount=goal,
        progress_pct=pct,
        donor_count=int(donation_stats[2] or 0),
        last_donation_at=donation_stats[3],
    )


def _fund_detail_out(
    fund: models.Fund,
    stats: schemas.FundStats,
) -> schemas.FundDetailOut:
    return schemas.FundDetailOut(
        id=fund.id,
        name=fund.name,
        description=fund.description,
        goal_amount=fund.goal_amount,
        start_date=fund.start_date,
        expected_end_date=fund.expected_end_date,
        completed_at=fund.completed_at,
        status=fund.status,
        is_active=fund.is_active,
        is_archived=fund.is_archived,
        created_by=fund.created_by,
        created_by_id=fund.created_by_id,
        created_at=fund.created_at,
        updated_at=fund.updated_at,
        archived_at=fund.archived_at,
        archived_by=fund.archived_by,
        stats=stats,
    )


def _batch_fund_stats(db: Session, fund_ids: list[int]) -> dict[int, schemas.FundStats]:
    """Single-pass SQL aggregation for donation and expense stats across multiple funds."""
    if not fund_ids:
        return {}

    don_rows = (
        db.query(
            models.Donation.fund_id,
            func.count(models.Donation.id),
            func.coalesce(func.sum(models.Donation.amount), 0.0),
            func.count(func.distinct(models.Donation.phone)),
            func.max(models.Donation.created_at),
        )
        .filter(models.Donation.fund_id.in_(fund_ids))
        .group_by(models.Donation.fund_id)
        .all()
    )
    don_map = {r[0]: r for r in don_rows}

    exp_rows = (
        db.query(
            models.Expense.fund_id,
            func.count(models.Expense.id),
            func.coalesce(func.sum(models.Expense.amount), 0.0),
        )
        .filter(
            models.Expense.fund_id.in_(fund_ids),
            models.Expense.is_deleted == False,
            models.Expense.approved_at != None,   # only approved expenses count
        )
        .group_by(models.Expense.fund_id)
        .all()
    )
    exp_map = {r[0]: r for r in exp_rows}

    # Pre-fetch goals in one query
    funds_map = {f.id: f for f in db.query(models.Fund).filter(models.Fund.id.in_(fund_ids)).all()}

    result: dict[int, schemas.FundStats] = {}
    for fid in fund_ids:
        dr = don_map.get(fid)
        er = exp_map.get(fid)
        total_collected = round(float(dr[2]) if dr else 0.0, 2)
        total_spent     = round(float(er[2]) if er else 0.0, 2)
        balance         = round(total_collected - total_spent, 2)
        goal = funds_map[fid].goal_amount if fid in funds_map else None
        pct  = round(total_collected / goal * 100, 1) if goal and goal > 0 else None
        result[fid] = schemas.FundStats(
            total_donations=int(dr[1]) if dr else 0,
            total_collected=total_collected,
            total_expenses=int(er[1]) if er else 0,
            total_spent=total_spent,
            balance=balance,
            goal_amount=goal,
            progress_pct=pct,
            donor_count=int(dr[3]) if dr else 0,
            last_donation_at=dr[4] if dr else None,
        )
    return result


# ─────────────────────────────────────────────
# DASHBOARD
# ─────────────────────────────────────────────
@router.get("/dashboard", response_model=schemas.FundDashboard)
def fund_dashboard(
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    funds = db.query(models.Fund).order_by(models.Fund.created_at.desc()).all()
    fund_ids = [f.id for f in funds]

    # Single-pass batch aggregation — no N+1
    stats_map = _batch_fund_stats(db, fund_ids)

    details = [_fund_detail_out(f, stats_map.get(f.id, _empty_stats(f))) for f in funds]

    total_collected = sum(s.total_collected for s in stats_map.values())
    total_spent     = sum(s.total_spent for s in stats_map.values())
    active_count    = sum(1 for f in funds if f.status == "active" and not f.is_archived)

    return schemas.FundDashboard(
        total_funds=sum(1 for f in funds if not f.is_archived),
        active_funds=active_count,
        total_collected=round(total_collected, 2),
        total_spent=round(total_spent, 2),
        overall_balance=round(total_collected - total_spent, 2),
        funds=details,
    )


def _empty_stats(fund: models.Fund) -> schemas.FundStats:
    return schemas.FundStats(
        total_donations=0, total_collected=0.0,
        total_expenses=0, total_spent=0.0,
        balance=0.0, goal_amount=fund.goal_amount,
        progress_pct=None, donor_count=0, last_donation_at=None,
    )


# ─────────────────────────────────────────────
# PUBLIC — active funds visible to all logged-in users
# ─────────────────────────────────────────────
@router.get("/public")
def list_funds_public(
    db: Session = Depends(get_db),
    user: dict = Depends(get_current_user),
):
    funds = (
        db.query(models.Fund)
        .filter(models.Fund.status == "active", models.Fund.is_archived == False)
        .order_by(models.Fund.created_at.desc())
        .all()
    )
    return [
        {
            "id": f.id,
            "name": f.name,
            "description": f.description or "",
            "goal_amount": f.goal_amount,
        }
        for f in funds
    ]


# ─────────────────────────────────────────────
# LIST FUNDS
# ─────────────────────────────────────────────
@router.get("/", response_model=List[schemas.FundDetailOut])
def list_funds(
    status: Optional[str] = None,
    include_archived: bool = False,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    query = db.query(models.Fund)
    if not include_archived:
        query = query.filter(models.Fund.is_archived == False)
    if status:
        query = query.filter(models.Fund.status == status)
    funds = query.order_by(models.Fund.created_at.desc()).all()
    stats_map = _batch_fund_stats(db, [f.id for f in funds])
    return [_fund_detail_out(f, stats_map.get(f.id, _empty_stats(f))) for f in funds]


# ─────────────────────────────────────────────
# CREATE FUND
# ─────────────────────────────────────────────
@router.post("/", response_model=schemas.FundOut, status_code=201)
def create_fund(
    data: schemas.FundCreate,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    uid, uname = _actor(user, db)

    requested_status = data.status or "active"
    resolved_status = _LEGACY_STATUS_MAP.get(requested_status, requested_status)
    if resolved_status not in VALID_STATUSES:
        raise HTTPException(422, "status must be one of: active, archived")

    if db.query(models.Fund).filter_by(name=data.name).first():
        raise HTTPException(409, "A fund with this name already exists")

    fund = models.Fund(
        name=data.name,
        description=data.description,
        goal_amount=data.goal_amount,
        start_date=data.start_date,
        expected_end_date=data.expected_end_date,
        status=resolved_status,
        created_by=uname,
        created_by_id=uid,
    )
    db.add(fund)
    db.flush()

    write_audit(db, "funds", fund.id, "create",
                new_values={"name": fund.name, "goal_amount": fund.goal_amount},
                performed_by_id=uid)
    db.commit()
    db.refresh(fund)

    invalidate_dashboard_cache()
    manager.publish_sync("finance", "fund_created", {"fund_id": fund.id, "name": fund.name})
    manager.publish_sync("finance", "dashboard_updated", {})
    return fund


# ─────────────────────────────────────────────
# GET SINGLE FUND
# ─────────────────────────────────────────────
@router.get("/{fund_id}", response_model=schemas.FundDetailOut)
def get_fund(
    fund_id: int,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    fund = _get_fund(db, fund_id)
    stats = _fund_stats(db, fund_id)
    return _fund_detail_out(fund, stats)


# ─────────────────────────────────────────────
# UPDATE FUND
# ─────────────────────────────────────────────
@router.put("/{fund_id}", response_model=schemas.FundOut)
def update_fund(
    fund_id: int,
    data: schemas.FundUpdate,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    uid, _ = _actor(user, db)
    fund = _get_fund(db, fund_id)

    if fund.is_archived:
        raise HTTPException(400, "Cannot update an archived fund")

    if data.status is not None:
        resolved = _LEGACY_STATUS_MAP.get(data.status, data.status)
        if resolved not in VALID_STATUSES:
            raise HTTPException(422, "status must be one of: active, archived")

    old = {"name": fund.name, "status": fund.status, "goal_amount": fund.goal_amount}

    if data.name is not None:
        existing = db.query(models.Fund).filter(models.Fund.name == data.name, models.Fund.id != fund_id).first()
        if existing:
            raise HTTPException(409, "Another fund with this name already exists")
        fund.name = data.name
    if data.description is not None:
        fund.description = data.description
    if data.goal_amount is not None:
        fund.goal_amount = data.goal_amount
    if data.start_date is not None:
        fund.start_date = data.start_date
    if data.expected_end_date is not None:
        fund.expected_end_date = data.expected_end_date
    if data.status is not None:
        fund.status = _LEGACY_STATUS_MAP.get(data.status, data.status)
    if data.is_active is not None:
        fund.is_active = data.is_active

    write_audit(db, "funds", fund.id, "update",
                old_values=old,
                new_values={"name": fund.name, "status": fund.status},
                performed_by_id=uid)
    db.commit()
    db.refresh(fund)

    invalidate_dashboard_cache()
    manager.publish_sync("finance", "fund_updated", {"fund_id": fund.id})
    manager.publish_sync("finance", "dashboard_updated", {})
    return fund


# ─────────────────────────────────────────────
# ARCHIVE FUND
# ─────────────────────────────────────────────
@router.patch("/{fund_id}/archive")
def archive_fund(
    fund_id: int,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    uid, uname = _actor(user, db)
    fund = _get_fund(db, fund_id)

    if fund.is_archived:
        raise HTTPException(400, "Fund is already archived")

    fund.is_archived = True
    fund.is_active   = False
    fund.status      = "archived"
    fund.archived_at = datetime.utcnow()
    fund.archived_by = uname

    write_audit(db, "funds", fund.id, "archive", performed_by_id=uid)
    db.commit()

    invalidate_dashboard_cache()
    manager.publish_sync("finance", "fund_archived", {"fund_id": fund.id})
    manager.publish_sync("finance", "dashboard_updated", {})
    return {"message": "Fund archived", "fund_id": fund.id}


# ─────────────────────────────────────────────
# UNARCHIVE FUND
# ─────────────────────────────────────────────
@router.patch("/{fund_id}/unarchive")
def unarchive_fund(
    fund_id: int,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    uid, _ = _actor(user, db)
    fund = _get_fund(db, fund_id)

    if not fund.is_archived:
        raise HTTPException(400, "Fund is not archived")

    fund.is_archived = False
    fund.is_active   = True
    fund.status      = "active"
    fund.archived_at = None
    fund.archived_by = None

    write_audit(db, "funds", fund.id, "unarchive", performed_by_id=uid)
    db.commit()
    invalidate_dashboard_cache()
    manager.publish_sync("finance", "fund_unarchived", {"fund_id": fund.id})
    manager.publish_sync("finance", "dashboard_updated", {})
    return {"message": "Fund unarchived", "fund_id": fund.id}


# ─────────────────────────────────────────────
# DELETE FUND (only if no transactions)
# ─────────────────────────────────────────────
@router.delete("/{fund_id}")
def delete_fund(
    fund_id: int,
    db: Session = Depends(get_db),
    user: dict = Depends(require_superadmin),
):
    uid, _ = _actor(user, db)
    fund = _get_fund(db, fund_id)

    has_donations = db.query(models.Donation).filter_by(fund_id=fund_id).count() > 0
    has_expenses  = db.query(models.Expense).filter(
        models.Expense.fund_id == fund_id,
        models.Expense.is_deleted == False,
    ).count() > 0

    if has_donations or has_expenses:
        raise HTTPException(
            409,
            "Cannot delete a fund that has donations or expenses. Archive it instead."
        )

    write_audit(db, "funds", fund.id, "delete",
                old_values={"name": fund.name},
                performed_by_id=uid)
    db.delete(fund)
    db.commit()

    invalidate_dashboard_cache()
    manager.publish_sync("finance", "fund_deleted", {"fund_id": fund_id})
    manager.publish_sync("finance", "dashboard_updated", {})
    return {"message": "Fund deleted", "fund_id": fund_id}


# ─────────────────────────────────────────────
# FUND COLLECTIONS (paginated)
# ─────────────────────────────────────────────
@router.get("/{fund_id}/collections")
def fund_collections(
    fund_id: int,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: Optional[str] = None,
    method: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    _get_fund(db, fund_id)

    query = (
        db.query(models.Donation)
        .options(joinedload(models.Donation.fund_rel))
        .filter(models.Donation.fund_id == fund_id)
    )
    if search:
        term = f"%{search}%"
        query = query.filter(or_(
            models.Donation.donor_name.ilike(term),
            models.Donation.receipt_id.ilike(term),
            models.Donation.phone.ilike(term),
            models.Donation.chanda_no.ilike(term),
        ))
    if method:
        query = query.filter(models.Donation.method == method)
    if from_date:
        query = query.filter(models.Donation.created_at >= datetime.fromisoformat(from_date))
    if to_date:
        query = query.filter(models.Donation.created_at <= datetime.fromisoformat(to_date))

    total = query.count()
    items = query.order_by(models.Donation.created_at.desc()).offset((page - 1) * page_size).limit(page_size).all()

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": max(1, math.ceil(total / page_size)),
        "items": [schemas.DonationOut(
            id=d.id,
            donor_name=d.donor_name,
            amount=d.amount,
            method=d.method,
            note=d.note,
            recorded_by=d.recorded_by,
            created_at=d.created_at,
            head_id=d.head_id,
            receipt_id=d.receipt_id,
            purpose_id=d.purpose_id,
            fund_id=d.fund_id,
            fund_name=fund_id and d.fund_rel and d.fund_rel.name,
            donor_type=d.donor_type,
            phone=d.phone,
            chanda_no=d.chanda_no,
            donation_date=d.donation_date,
            receipt_image=d.receipt_image,
        ) for d in items],
    }


# ─────────────────────────────────────────────
# FUND EXPENSES (paginated)
# ─────────────────────────────────────────────
@router.get("/{fund_id}/expenses")
def fund_expenses(
    fund_id: int,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    _get_fund(db, fund_id)

    query = (
        db.query(models.Expense)
        .options(joinedload(models.Expense.category_rel))
        .filter(
            models.Expense.fund_id == fund_id,
            models.Expense.is_deleted == False,
            models.Expense.approved_at != None,   # only approved expenses shown in fund
        )
    )
    if search:
        term = f"%{search}%"
        query = query.filter(or_(
            models.Expense.title.ilike(term),
            models.Expense.receipt_id.ilike(term),
            models.Expense.vendor_name.ilike(term),
        ))
    if from_date:
        query = query.filter(models.Expense.created_at >= datetime.fromisoformat(from_date))
    if to_date:
        query = query.filter(models.Expense.created_at <= datetime.fromisoformat(to_date))

    total = query.count()
    items = query.order_by(models.Expense.created_at.desc()).offset((page - 1) * page_size).limit(page_size).all()

    def _build(e: models.Expense) -> schemas.ExpenseOut:
        cat_name = e.category_rel.name if e.category_rel else e.category
        return schemas.ExpenseOut(
            id=e.id,
            title=e.title,
            amount=e.amount,
            category=e.category,
            category_id=e.category_id,
            category_name=cat_name,
            fund_id=e.fund_id,
            vendor_name=e.vendor_name,
            expense_date=e.expense_date,
            note=e.note,
            receipt_image=e.receipt_image,
            receipt_id=e.receipt_id,
            created_by=e.created_by,
            created_by_id=e.created_by_id,
            approved_by=e.approved_by,
            approved_by_id=e.approved_by_id,
            approved_at=e.approved_at,
            created_at=e.created_at,
            is_deleted=e.is_deleted,
        )

    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": max(1, math.ceil(total / page_size)),
        "items": [_build(e) for e in items],
    }


# ─────────────────────────────────────────────
# FUND STATS
# ─────────────────────────────────────────────
@router.get("/{fund_id}/stats", response_model=schemas.FundStats)
def get_fund_stats(
    fund_id: int,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    _get_fund(db, fund_id)
    return _fund_stats(db, fund_id)


# ─────────────────────────────────────────────
# FUND REPORT
# ─────────────────────────────────────────────
@router.get("/{fund_id}/report")
def fund_report(
    fund_id: int,
    format: str = Query("pdf", pattern="^(pdf|excel)$"),
    from_date: Optional[str] = Query(None, description="ISO date e.g. 2026-01-01"),
    to_date: Optional[str] = Query(None, description="ISO date e.g. 2026-12-31"),
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    fund  = _get_fund(db, fund_id)
    stats = _fund_stats(db, fund_id)

    don_q = (
        db.query(models.Donation)
        .filter(models.Donation.fund_id == fund_id)
    )
    exp_q = (
        db.query(models.Expense)
        .options(joinedload(models.Expense.category_rel))
        .filter(models.Expense.fund_id == fund_id, models.Expense.is_deleted == False)
    )
    if from_date:
        _from = datetime.fromisoformat(from_date)
        don_q = don_q.filter(models.Donation.created_at >= _from)
        exp_q = exp_q.filter(models.Expense.created_at >= _from)
    if to_date:
        _to = datetime.fromisoformat(to_date)
        don_q = don_q.filter(models.Donation.created_at <= _to)
        exp_q = exp_q.filter(models.Expense.created_at <= _to)

    donations = don_q.order_by(models.Donation.created_at.desc()).all()
    expenses  = exp_q.order_by(models.Expense.created_at.desc()).all()

    if format == "excel":
        return _excel_report(fund, stats, donations, expenses)
    else:
        return _pdf_report(fund, stats, donations, expenses)


def _csv_report(fund, stats, donations, expenses):
    import csv

    buf = io.StringIO()
    w   = csv.writer(buf)

    w.writerow(["FUND REPORT", fund.name])
    w.writerow(["Generated", datetime.utcnow().strftime("%Y-%m-%d %H:%M")])
    w.writerow([])
    w.writerow(["--- SUMMARY ---"])
    w.writerow(["Total Collected", stats.total_collected])
    w.writerow(["Total Spent",     stats.total_spent])
    w.writerow(["Balance",         stats.balance])
    w.writerow(["Goal Amount",     stats.goal_amount or "N/A"])
    w.writerow(["Progress",        f"{stats.progress_pct}%" if stats.progress_pct else "N/A"])
    w.writerow([])
    w.writerow(["--- DONATIONS ---"])
    w.writerow(["Receipt ID", "Date", "Donor", "Phone", "Chanda No", "Method", "Amount", "Note"])
    for d in donations:
        w.writerow([
            d.receipt_id or "",
            d.created_at.strftime("%Y-%m-%d") if d.created_at else "",
            d.donor_name,
            d.phone or "",
            d.chanda_no or "",
            d.method,
            d.amount,
            d.note or "",
        ])
    w.writerow([])
    w.writerow(["--- EXPENSES ---"])
    w.writerow(["Receipt ID", "Date", "Title", "Category", "Vendor", "Amount", "Note"])
    for e in expenses:
        cat_name = e.category_rel.name if e.category_rel else (e.category or "")
        w.writerow([
            e.receipt_id or "",
            e.created_at.strftime("%Y-%m-%d") if e.created_at else "",
            e.title,
            cat_name,
            e.vendor_name or "",
            e.amount,
            e.note or "",
        ])

    content = buf.getvalue().encode("utf-8-sig")
    fname   = f"fund_{fund.id}_{fund.name.replace(' ', '_')}.csv"
    return StreamingResponse(
        io.BytesIO(content),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


def _excel_report(fund, stats, donations, expenses):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")

    wb = openpyxl.Workbook()

    # ── Summary sheet ──
    ws = wb.active
    ws.title = "Summary"
    hdr_fill = PatternFill("solid", fgColor="1E3A5F")
    hdr_font = Font(bold=True, color="FFFFFF")

    def _hdr(cell, val):
        cell.value = val
        cell.font  = hdr_font
        cell.fill  = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    _hdr(ws["A1"], f"Fund Report: {fund.name}")
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws.merge_cells("A1:B1")
    ws["A2"] = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')}"
    ws["A4"] = "Total Collected";  ws["B4"] = stats.total_collected
    ws["A5"] = "Total Spent";      ws["B5"] = stats.total_spent
    ws["A6"] = "Balance";          ws["B6"] = stats.balance
    ws["A7"] = "Goal Amount";      ws["B7"] = stats.goal_amount or "N/A"
    ws["A8"] = "Progress";         ws["B8"] = f"{stats.progress_pct}%" if stats.progress_pct else "N/A"
    ws["A9"] = "Donors";           ws["B9"] = stats.donor_count

    # ── Donations sheet ──
    wd = wb.create_sheet("Donations")
    headers = ["Receipt ID", "Date", "Donor", "Phone", "Chanda No", "Method", "Amount", "Donor Type", "Note"]
    for col, h in enumerate(headers, 1):
        _hdr(wd.cell(1, col), h)
    for row, d in enumerate(donations, 2):
        wd.cell(row, 1).value = d.receipt_id or ""
        wd.cell(row, 2).value = d.created_at.strftime("%Y-%m-%d") if d.created_at else ""
        wd.cell(row, 3).value = d.donor_name
        wd.cell(row, 4).value = d.phone or ""
        wd.cell(row, 5).value = d.chanda_no or ""
        wd.cell(row, 6).value = d.method
        wd.cell(row, 7).value = d.amount
        wd.cell(row, 8).value = d.donor_type or ""
        wd.cell(row, 9).value = d.note or ""
    for col in range(1, len(headers) + 1):
        wd.column_dimensions[get_column_letter(col)].auto_size = True

    # ── Expenses sheet ──
    we = wb.create_sheet("Expenses")
    eheaders = ["Receipt ID", "Date", "Title", "Category", "Vendor", "Amount", "Note"]
    for col, h in enumerate(eheaders, 1):
        _hdr(we.cell(1, col), h)
    for row, e in enumerate(expenses, 2):
        cat_name = e.category_rel.name if e.category_rel else (e.category or "")
        we.cell(row, 1).value = e.receipt_id or ""
        we.cell(row, 2).value = e.created_at.strftime("%Y-%m-%d") if e.created_at else ""
        we.cell(row, 3).value = e.title
        we.cell(row, 4).value = cat_name
        we.cell(row, 5).value = e.vendor_name or ""
        we.cell(row, 6).value = e.amount
        we.cell(row, 7).value = e.note or ""

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"fund_{fund.id}_{fund.name.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


def _pdf_report(fund, stats, donations, expenses):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
        )
    except ImportError:
        raise HTTPException(500, "reportlab not installed")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    story  = []

    DARK   = colors.HexColor("#1E3A5F")
    LIGHT  = colors.HexColor("#EBF0F7")

    story.append(Paragraph(f"Fund Report: {fund.name}", styles["Title"]))
    story.append(Paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC", styles["Normal"]))
    story.append(Spacer(1, 0.5*cm))

    # Summary table
    summary_data = [
        ["Metric",           "Value"],
        ["Total Collected",  f"₹{stats.total_collected:,.2f}"],
        ["Total Spent",      f"₹{stats.total_spent:,.2f}"],
        ["Balance",          f"₹{stats.balance:,.2f}"],
        ["Goal Amount",      f"₹{stats.goal_amount:,.2f}" if stats.goal_amount else "N/A"],
        ["Progress",         f"{stats.progress_pct}%" if stats.progress_pct else "N/A"],
        ["Total Donors",     str(stats.donor_count)],
        ["Total Donations",  str(stats.total_donations)],
        ["Total Expenses",   str(stats.total_expenses)],
    ]
    summary_table = Table(summary_data, colWidths=[8*cm, 8*cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), DARK),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 1), (-1, -1), LIGHT),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("GRID",       (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE",   (0, 0), (-1, -1), 10),
        ("PADDING",    (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 0.5*cm))

    # Donations table
    story.append(Paragraph("Donations", styles["Heading2"]))
    don_headers = ["Date", "Donor", "Phone", "Method", "Amount", "Receipt"]
    don_rows    = [don_headers]
    for d in donations:
        don_rows.append([
            d.created_at.strftime("%Y-%m-%d") if d.created_at else "",
            d.donor_name[:25],
            d.phone or "",
            d.method,
            f"₹{d.amount:,.2f}",
            d.receipt_id or "",
        ])
    if len(don_rows) > 1:
        don_table = Table(don_rows, colWidths=[2.5*cm, 4*cm, 3*cm, 2*cm, 2.5*cm, 3*cm])
        don_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), DARK),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("GRID",       (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ("FONTSIZE",   (0, 0), (-1, -1), 8),
            ("PADDING",    (0, 0), (-1, -1), 4),
        ]))
        story.append(don_table)
    else:
        story.append(Paragraph("No donations recorded.", styles["Normal"]))

    story.append(Spacer(1, 0.5*cm))

    # Expenses table
    story.append(Paragraph("Expenses", styles["Heading2"]))
    exp_headers = ["Date", "Title", "Category", "Vendor", "Amount", "Receipt"]
    exp_rows    = [exp_headers]
    for e in expenses:
        cat_name = e.category_rel.name if e.category_rel else (e.category or "—")
        exp_rows.append([
            e.created_at.strftime("%Y-%m-%d") if e.created_at else "",
            e.title[:25],
            cat_name[:15],
            (e.vendor_name or "")[:15],
            f"₹{e.amount:,.2f}",
            e.receipt_id or "",
        ])
    if len(exp_rows) > 1:
        exp_table = Table(exp_rows, colWidths=[2.5*cm, 4*cm, 2.5*cm, 3*cm, 2.5*cm, 2.5*cm])
        exp_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), DARK),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("GRID",       (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ("FONTSIZE",   (0, 0), (-1, -1), 8),
            ("PADDING",    (0, 0), (-1, -1), 4),
        ]))
        story.append(exp_table)
    else:
        story.append(Paragraph("No expenses recorded.", styles["Normal"]))

    doc.build(story)
    buf.seek(0)
    fname = f"fund_{fund.id}_{fund.name.replace(' ', '_')}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
